from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate


ROWS = [
   {"assignment": "Design Review", "course": "CS 491", "priority": "High", "due_date": "2026-09-30"},
   {"assignment": "Lab 2", "course": "CS 491", "priority": "Medium", "due_date": "2026-10-02"},
]


def create_docxtpl_template(template_path: Path):
   template = Document()
   template.add_heading("{{ title }}", level=0)
   template.add_paragraph("Student: {{ student }}")
   template.add_paragraph("Project: {{ project }}")
   template.add_paragraph("Assignments:")
   table = template.add_table(rows=1, cols=4)
   table.rows[0].cells[0].text = "Assignment"
   table.rows[0].cells[1].text = "Course"
   table.rows[0].cells[2].text = "Priority"
   table.rows[0].cells[3].text = "Due"
   row = table.add_row().cells
   row[0].text = "{{ rows[0].assignment }}"
   row[1].text = "{{ rows[0].course }}"
   row[2].text = "{{ rows[0].priority }}"
   row[3].text = "{{ rows[0].due_date }}"
   template.save(template_path)


def python_docx_candidate(path: Path):
   doc = Document()
   doc.add_heading("Folio Assignment Summary", 0)
   doc.add_paragraph("Student: Alex Rivera")
   doc.add_paragraph("Project: Senior Design")
   doc.add_heading("Assignments", level=1)
   table = doc.add_table(rows=1, cols=4)
   table.rows[0].cells[0].text = "Assignment"
   table.rows[0].cells[1].text = "Course"
   table.rows[0].cells[2].text = "Priority"
   table.rows[0].cells[3].text = "Due"
   for row in ROWS:
      cells = table.add_row().cells
      cells[0].text = row["assignment"]
      cells[1].text = row["course"]
      cells[2].text = row["priority"]
      cells[3].text = row["due_date"]
   doc.save(path)
   reloaded = Document(path)
   paras = [p.text for p in reloaded.paragraphs]
   assert "Folio Assignment Summary" in paras
   assert any("Design Review" in p for p in [cell.text for row in reloaded.tables[0].rows for cell in row.cells])
   return path


def docxtpl_candidate(path: Path, template_path: Path):
   tpl = DocxTemplate(str(template_path))
   context = {"title": "Folio Assignment Summary", "student": "Alex Rivera", "project": "Senior Design", "rows": ROWS}
   tpl.render(context)
   tpl.save(path)
   assert path.exists() and path.stat().st_size > 0
   return path


def test_docx_candidates(tmp_path):
   template_path = tmp_path / "template_report.docx"
   create_docxtpl_template(template_path)
   python_docx = python_docx_candidate(tmp_path / "folio_python_docx.docx")
   docxtpl_path = tmp_path / "folio_docxtpl.docx"
   docxtpl_candidate(docxtpl_path, template_path)
   assert python_docx.exists()
   assert docxtpl_path.exists()


def test_docx_repeated_generation(tmp_path):
   for idx in range(2):
      python_docx_candidate(tmp_path / f"repeat_{idx}.docx")
